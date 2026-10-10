import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_OUT = Path("docs/reports/BAO_CAO_DO_AN_TRAFFICGUARD_90_TRANG.docx")
ASSETS_DIR = Path("report_assets")
DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_styled_document():
    doc = Document()
    
    # Page Setup: Left 3cm, Right 2cm, Top 2cm, Bottom 2cm
    for section in doc.sections:
        section.top_margin = Inches(0.79)      # 2.0 cm
        section.bottom_margin = Inches(0.79)   # 2.0 cm
        section.left_margin = Inches(1.18)     # 3.0 cm
        section.right_margin = Inches(0.79)    # 2.0 cm
        
        # Header setup
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("TRAFFICGUARD | BÁO CÁO ĐỒ ÁN TỐT NGHIỆP HỆ THỐNG TRỢ LÝ LÁI XE AN TOÀN")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)
        hrun.font.italic = True
        
        # Footer setup
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Hệ Thống Trợ Lý Lái Xe An Toàn & Giám Sát Đội Xe Doanh Nghiệp — TrafficGuard")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)
        
    return doc

def add_chapter_heading(doc, chapter_num, title):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    
    r_ch = p.add_run(f"CHƯƠNG {chapter_num}\n")
    r_ch.font.name = "Times New Roman"
    r_ch.font.size = Pt(14)
    r_ch.font.bold = True
    r_ch.font.color.rgb = RGBColor(30, 58, 138)
    
    r_ti = p.add_run(title.upper())
    r_ti.font.name = "Times New Roman"
    r_ti.font.size = Pt(16)
    r_ti.font.bold = True
    r_ti.font.color.rgb = RGBColor(15, 44, 89)

def add_section_heading(doc, sec_num, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(f"{sec_num}. {title}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138)

def add_subsection_heading(doc, subsec_num, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(f"{subsec_num}. {title}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)

def add_paragraph(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
        
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(11.5)
    r_text.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_bullet_point(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.4)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
        
    r_text = p.add_run(text)
    r_text.font.name = "Times New Roman"
    r_text.font.size = Pt(11.5)
    r_text.font.color.rgb = RGBColor(15, 23, 42)

def add_callout_box(doc, text, title="LƯU Ý KỸ THUẬT / CĂN CỨ PHÁP LÝ"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(2)
    
    r_ti = p.add_run(f"📌 {title}: ")
    r_ti.font.name = "Times New Roman"
    r_ti.font.size = Pt(11)
    r_ti.font.bold = True
    r_ti.font.color.rgb = RGBColor(30, 58, 138)
    
    r_tx = p.add_run(text)
    r_tx.font.name = "Times New Roman"
    r_tx.font.size = Pt(11)
    r_tx.font.italic = True
    r_tx.font.color.rgb = RGBColor(30, 41, 59)
    
    # Space after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

def add_formula_box(doc, formula_str, formula_num, explanation=None):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell_f = table.rows[0].cells[0]
    cell_n = table.rows[0].cells[1]
    
    cell_f.width = Inches(5.6)
    cell_n.width = Inches(0.9)
    
    set_cell_background(cell_f, "F8FAFC")
    set_cell_background(cell_n, "F8FAFC")
    set_cell_margins(cell_f, 100, 100, 140, 140)
    set_cell_margins(cell_n, 100, 100, 100, 100)
    
    p_f = cell_f.paragraphs[0]
    p_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f = p_f.add_run(formula_str)
    r_f.font.name = "Times New Roman"
    r_f.font.size = Pt(12)
    r_f.font.bold = True
    r_f.font.color.rgb = RGBColor(15, 23, 42)
    
    p_n = cell_n.paragraphs[0]
    p_n.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_n = p_n.add_run(f"({formula_num})")
    r_n.font.name = "Times New Roman"
    r_n.font.size = Pt(11)
    r_n.font.bold = True
    r_n.font.color.rgb = RGBColor(71, 85, 105)
    
    if explanation:
        p_ex = doc.add_paragraph()
        p_ex.paragraph_format.line_spacing = 1.2
        p_ex.paragraph_format.space_after = Pt(6)
        r_ex = p_ex.add_run(f"Trong đó: {explanation}")
        r_ex.font.name = "Times New Roman"
        r_ex.font.size = Pt(10.5)
        r_ex.font.italic = True
        r_ex.font.color.rgb = RGBColor(71, 85, 105)

def add_custom_table(doc, headers, data, col_widths=None, table_title=""):
    if table_title:
        p_ti = doc.add_paragraph()
        p_ti.paragraph_format.space_before = Pt(8)
        p_ti.paragraph_format.space_after = Pt(4)
        p_ti.paragraph_format.keep_with_next = True
        r_ti = p_ti.add_run(table_title)
        r_ti.font.name = "Times New Roman"
        r_ti.font.size = Pt(11)
        r_ti.font.bold = True
        r_ti.font.color.rgb = RGBColor(30, 41, 59)
        
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Headers
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], 120, 120, 140, 140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        if col_widths and i < len(col_widths):
            hdr_cells[i].width = col_widths[i]
            
    # Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_background(row_cells[col_idx], bg_col)
            set_cell_margins(row_cells[col_idx], 80, 80, 120, 120)
            p = row_cells[col_idx].paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
                if col_idx == 0:
                    run.font.bold = True
            if col_widths and col_idx < len(col_widths):
                row_cells[col_idx].width = col_widths[col_idx]
                
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

def add_image_figure(doc, img_path, caption, width_in=6.0):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = True
    
    if Path(img_path).exists():
        run_img = p_img.add_run()
        run_img.add_picture(str(img_path), width=Inches(width_in))
    else:
        p_img.add_run(f"[HÌNH ẢNH MINH HỌA: {img_path}]").font.italic = True
        
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_after = Pt(12)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10.5)
    r_cap.font.bold = True
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(71, 85, 105)
