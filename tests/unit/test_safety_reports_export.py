import pytest
import pytest_asyncio
import openpyxl
import io
from sqlalchemy.ext.asyncio import create_async_engine, AsyncSession
from sqlalchemy.orm import sessionmaker

from backend.app.database.base import Base
from backend.app.models_db.organization import Organization
from backend.app.models_db.driver import Driver
from backend.app.models_db.vehicle import Vehicle
from backend.app.models_db.journey import Journey
from backend.app.models_db.violation import Violation
from backend.app.services.report_generator import (
    generate_safety_summary,
    export_safety_report_excel,
    export_safety_report_pdf
)


@pytest_asyncio.fixture
async def async_db():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async_session = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    async with async_session() as session:
        org = Organization(id=1, name="Express Transport")
        session.add(org)

        driver1 = Driver(
            id=1,
            company_id=1,
            email="driver1@express.vn",
            full_name="Nguyễn Văn A",
            role="FLEET_DRIVER",
            password_hash="pwd",
            safety_score=65.0  # Rủi ro cao (< 80)
        )
        driver2 = Driver(
            id=2,
            company_id=1,
            email="driver2@express.vn",
            full_name="Trần Văn B",
            role="FLEET_DRIVER",
            password_hash="pwd",
            safety_score=95.0  # An toàn
        )
        session.add_all([driver1, driver2])

        # Tạo một số chuyến đi
        j1 = Journey(id=1, company_id=1, journey_code="J1", vehicle_id=1, driver_id=1, status="COMPLETED")
        j2 = Journey(id=2, company_id=1, journey_code="J2", vehicle_id=1, driver_id=2, status="ACTIVE")
        session.add_all([j1, j2])

        # Tạo vi phạm
        v1 = Violation(
            id=1,
            company_id=1,
            driver_id=1,
            violation_type="drowsy",
            severity="CRITICAL",
            review_status="CONFIRMED"
        )
        v2 = Violation(
            id=2,
            company_id=1,
            driver_id=1,
            violation_type="speed_limit",
            severity="WARNING",
            review_status="FALSE_POSITIVE"
        )
        session.add_all([v1, v2])

        await session.commit()
        yield session

    await engine.dispose()


@pytest.mark.asyncio
async def test_generate_safety_summary(async_db: AsyncSession):
    summary = await generate_safety_summary(db=async_db, company_id=1, days=30)

    assert summary["total_journeys"] == 2
    assert summary["completed_journeys"] == 1
    assert summary["active_journeys"] == 1
    assert summary["total_violations"] == 2
    assert summary["confirmed_violations"] == 1
    assert summary["false_positives"] == 1

    # Tỷ lệ chính xác AI: 1 confirmed / 2 reviewed = 50.0%
    assert summary["ai_precision_rate"] == 50.0

    # Nhận diện tài xế rủi ro cao: driver 1 có safety_score=65.0
    assert len(summary["high_risk_drivers"]) == 1
    assert summary["high_risk_drivers"][0]["driver_id"] == 1
    assert summary["high_risk_drivers"][0]["safety_score"] == 65.0


def test_export_excel_and_pdf():
    mock_summary = {
        "period_days": 30,
        "generated_at": "2026-09-28 14:00:00 UTC",
        "total_journeys": 45,
        "completed_journeys": 40,
        "active_journeys": 4,
        "cancelled_journeys": 1,
        "planned_journeys": 0,
        "total_violations": 12,
        "confirmed_violations": 10,
        "false_positives": 2,
        "pending_review": 0,
        "ai_precision_rate": 83.3,
        "module_breakdown": {"drowsy": 4, "speed_limit": 5, "no_seatbelt": 3},
        "severity_breakdown": {"CRITICAL": 4, "WARNING": 8, "INFO": 0},
        "avg_fleet_safety_score": 88.5,
        "high_risk_drivers": [
            {
                "driver_id": 1,
                "full_name": "Nguyễn Văn Rủi Ro",
                "email": "driver1@test.com",
                "safety_score": 58.0,
                "confirmed_violations_count": 4,
                "risk_level": "RẤT CAO"
            }
        ]
    }

    # 1. Test Excel export
    excel_bytes = export_safety_report_excel(mock_summary)
    assert len(excel_bytes) > 1000
    # Validate with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Tổng Quan An Toàn" in wb.sheetnames
    assert "Tài Xế Rủi Ro Cao" in wb.sheetnames

    # 2. Test PDF export
    pdf_bytes = export_safety_report_pdf(mock_summary)
    assert len(pdf_bytes) > 1000
    assert pdf_bytes.startswith(b"%PDF")
