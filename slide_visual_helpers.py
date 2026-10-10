# -*- coding: utf-8 -*-
"""
build_visual_minimal_slides.py
Generates an ultra-visual, low-text presentation deck matching 745545451-Slide-UTH.pptx.
Features:
- Large visual canvas (60-70% of slide area).
- Minimalist text (punchy keywords, metrics, formulas, no walls of text).
- Real project diagrams + dedicated replacement-ready frames with clear captions.
- Zero synthetic AI images; easy for user to drop in their own live demo screenshots.
"""

import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

# Design Colors matching UTH branding
COLOR_UTH_TEAL = RGBColor(0, 102, 102)        # Primary Teal #006666
COLOR_UTH_TEAL_DARK = RGBColor(10, 77, 77)    # Deep Teal #0A4D4D
COLOR_UTH_TEAL_LIGHT = RGBColor(240, 253, 250)# Light Teal Card Background #F0FDFA
COLOR_UTH_RED = RGBColor(192, 0, 0)           # Crimson / Accent Red #C00000
COLOR_DARK_TEXT = RGBColor(15, 23, 42)        # Dark Slate #0F172A
COLOR_BODY_TEXT = RGBColor(51, 65, 85)        # Slate 700 #334155
COLOR_MUTED_TEXT = RGBColor(100, 116, 139)    # Slate 500 #64748B
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_CARD_BG = RGBColor(248, 250, 252)       # Light Slate #F8FAFC
COLOR_CARD_BORDER = RGBColor(203, 213, 225)   # Border #CBD5E1
COLOR_TEAL_BORDER = RGBColor(94, 234, 212)   # Teal Border #5EEAD4
COLOR_PHOTO_BG = RGBColor(241, 245, 249)      # Slate 100 for Photo Placeholders #F1F5F9
COLOR_SUCCESS = RGBColor(16, 185, 129)        # Green #10B981
COLOR_WARNING = RGBColor(217, 119, 6)         # Amber #D97706
COLOR_DANGER = RGBColor(239, 68, 68)          # Red #EF4444

FONT_MAIN = "Times New Roman"

def add_header(slide, title_text, category_text, slide_num, total_slides=28):
    """Adds standard UTH header and slide number"""
    tag_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.24), Inches(11.5), Inches(0.28))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_text.upper()
    p_tag.font.name = FONT_MAIN
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_UTH_RED

    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.52), Inches(11.5), Inches(0.65))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_MAIN
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_UTH_TEAL

    num_box = slide.shapes.add_textbox(Inches(10.5), Inches(6.85), Inches(2.0), Inches(0.35))
    tf_num = num_box.text_frame
    tf_num.word_wrap = False
    p_num = tf_num.paragraphs[0]
    p_num.alignment = PP_ALIGN.RIGHT
    p_num.text = f"{slide_num} / {total_slides}"
    p_num.font.name = FONT_MAIN
    p_num.font.size = Pt(11)
    p_num.font.bold = True
    p_num.font.color.rgb = COLOR_UTH_TEAL

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
    """Draws a clean rectangular background card"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.0)
    return shape

def add_image_or_placeholder(slide, img_path, left, top, width, height, placeholder_title, placeholder_hint):
    """
    Inserts image if it exists.
    Also adds a neat visual caption frame below or behind, so user can easily replace it!
    If image does not exist, draws an elegant screenshot placeholder frame.
    """
    # Outer frame
    frame = add_card(slide, left, top, width, height, bg_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER)
    
    img_added = False
    if img_path and os.path.exists(img_path):
        # Insert image slightly inset inside frame
        inset_x = Inches(0.08)
        inset_y = Inches(0.08)
        caption_h = Inches(0.45)
        slide.shapes.add_picture(img_path, left + inset_x, top + inset_y, width - (inset_x * 2), height - caption_h - inset_y)
        img_added = True

        # Caption badge at bottom
        tb = slide.shapes.add_textbox(left + inset_x, top + height - caption_h, width - (inset_x * 2), caption_h - Inches(0.05))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = f"🖼️ {placeholder_title}"
        p.font.name = FONT_MAIN
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_UTH_TEAL

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = f"💡 {placeholder_hint}"
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(9)
        p2.font.italic = True
        p2.font.color.rgb = COLOR_MUTED_TEXT
    else:
        # Full placeholder design
        tb = slide.shapes.add_textbox(left + Inches(0.3), top + height * 0.28, width - Inches(0.6), height * 0.5)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        p1.text = "📷 KHUNG CHÈN HÌNH ẢNH / ẢNH CHỤP THỰC TẾ"
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_UTH_TEAL

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = f"【 {placeholder_title} 】"
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_UTH_RED

        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.CENTER
        p3.text = f"👉 Gợi ý: {placeholder_hint}\n(Nhấp chuột phải -> Change Picture hoặc dán ảnh chụp màn hình/video test vào đây)"
        p3.font.name = FONT_MAIN
        p3.font.size = Pt(11)
        p3.font.color.rgb = COLOR_MUTED_TEXT

    return frame

def add_side_card(slide, left, top, width, height, title, items, badge_title=None, badge_color=COLOR_UTH_TEAL):
    """
    Draws a minimalist key-points card beside the visual.
    Items is a list of tuples: (Bold Keyword, Short Note, Optional Value/Tag)
    """
    card = add_card(slide, left, top, width, height, bg_color=COLOR_WHITE, border_color=badge_color)
    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.18), width - Inches(0.4), height - Inches(0.36))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = badge_color

    for item in items:
        p1 = tf.add_paragraph()
        p1.space_before = Pt(8)
        
        # Keyword badge
        run_kw = p1.add_run()
        run_kw.text = f"• {item[0]}: "
        run_kw.font.name = FONT_MAIN
        run_kw.font.size = Pt(11)
        run_kw.font.bold = True
        run_kw.font.color.rgb = COLOR_DARK_TEXT

        # Short description
        run_desc = p1.add_run()
        run_desc.text = item[1]
        run_desc.font.name = FONT_MAIN
        run_desc.font.size = Pt(10.5)
        run_desc.font.bold = False
        run_desc.font.color.rgb = COLOR_BODY_TEXT

        # Value highlight if provided
        if len(item) > 2 and item[2]:
            run_val = p1.add_run()
            run_val.text = f"  [{item[2]}]"
            run_val.font.name = FONT_MAIN
            run_val.font.size = Pt(10.5)
            run_val.font.bold = True
            run_val.font.color.rgb = COLOR_UTH_RED

    return card

def add_kpi_strip(slide, left, top, width, height, kpis):
    """Draws horizontal KPI metric strips"""
    n = len(kpis)
    kw = (width - Inches(0.15) * (n - 1)) / n
    for i, (num, label, color) in enumerate(kpis):
        l = left + i * (kw + Inches(0.15))
        add_card(slide, l, top, kw, height, bg_color=COLOR_WHITE, border_color=color)
        tb = slide.shapes.add_textbox(l + Inches(0.08), top + Inches(0.06), kw - Inches(0.16), height - Inches(0.12))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        p1.text = num
        p1.font.name = FONT_MAIN
        p1.font.size = Pt(20)
        p1.font.bold = True
        p1.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.text = label
        p2.font.name = FONT_MAIN
        p2.font.size = Pt(9.5)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_DARK_TEXT

print("Visual layout generator helper functions initialized.")
